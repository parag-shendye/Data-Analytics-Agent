"""Word input/output for the agent: read questions from a .docx, write answers to a .docx.

Tables in the report are built from tool results (exact data), not from the
model's text, so their numbers are always correct.
"""

from datetime import datetime
from pathlib import Path

from docx import Document

from delinquency_agent.config import settings


# ---------------------------------------------------------------- reading

def read_questions(path: str) -> list[str]:
    """Paragraphs or first-column table cells in the .docx that end with '?'."""
    doc = Document(path)
    questions = [p.text.strip() for p in doc.paragraphs if p.text.strip().endswith("?")]
    for table in doc.tables:
        for row in table.rows:
            text = row.cells[0].text.strip()
            if text.endswith("?"):
                questions.append(text)
    return questions


# ---------------------------------------------------------------- writing

def write_report(results: list[dict], path: str | None = None) -> Path:
    """One section per question: the model's answer, then a table for each list of records in the tool results."""
    doc = Document()
    doc.add_heading("Delinquency analysis", level=0)

    for i, result in enumerate(results, start=1):
        doc.add_heading(f"{i}. {result['question']}", level=1)
        add_answer_text(doc, result["answer"])
        for tool_result in result["tool_results"]:
            for title, rows in find_tables(tool_result["data"]):
                add_table(doc, title, rows)

    out = Path(path) if path else settings.output_dir / f"answers_{datetime.now():%Y%m%d_%H%M%S}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def add_answer_text(doc: Document, text: str) -> None:
    """Convert the model's light markdown into Word paragraphs.

    Markdown tables are skipped: the real tables come from the tool results.
    """
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("|"):
            continue
        if line.startswith(("- ", "* ")):
            doc.add_paragraph(strip_markdown(line[2:]), style="List Bullet")
        elif line.startswith("**") and line.rstrip(":").endswith("**"):  # "**Caveats:**" -> bold line
            doc.add_paragraph().add_run(strip_markdown(line)).bold = True
        elif line.startswith("#"):
            doc.add_paragraph().add_run(strip_markdown(line.lstrip("# "))).bold = True
        else:
            doc.add_paragraph(strip_markdown(line))


def strip_markdown(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def find_tables(data: dict | None) -> list[tuple[str, list[dict]]]:
    """Every value in a tool result that is a list of records, e.g. columns_with_missing."""
    if not isinstance(data, dict):
        return []
    return [(key, value) for key, value in data.items()
            if isinstance(value, list) and value and all(isinstance(v, dict) for v in value)]


def add_table(doc: Document, title: str, rows: list[dict]) -> None:
    """A titled Word table with one column per key of the records."""
    doc.add_paragraph().add_run(title.replace("_", " ").capitalize()).italic = True
    columns = list(rows[0].keys())
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, columns):
        cell.text = name.replace("_", " ").capitalize()
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        for cell, name in zip(table.add_row().cells, columns):
            cell.text = str(row.get(name, ""))
